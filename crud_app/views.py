import docx
from docx import Document
from django.shortcuts import render, redirect, get_object_or_404
from django.http import HttpResponse
from .models import Item

# ==========================================
# 1. Выгрузка ВСЕХ тестов проекта в Word
# ==========================================
def export_tests_to_word(request):
    doc = Document()
    doc.add_heading('Отчёт по юнит-тестам Django (blog/tests.py)', level=1)
    
    # Читаем исходный код твоих тестов
    try:
        with open('blog/tests.py', 'r', encoding='utf-8') as f:
            tests_code = f.read()
    except FileNotFoundError:
        tests_code = "Файл blog/tests.py не найден."

    doc.add_heading('Код тестов проекта:', level=2)
    doc.add_paragraph(tests_code)

    # Отдаём файл как скачиваемый документ Word
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    response['Content-Disposition'] = 'attachment; filename="django_tests_report.docx"'
    
    doc.save(response)
    return response


# ==========================================
# 2. Заготовка CRUD (для следующей лабы)
# ==========================================
def item_list(request):
    items = Item.objects.all()
    return render(request, 'crud_app/item_list.html', {'items': items})

def item_create(request):
    if request.method == 'POST':
        title = request.POST.get('title')
        content = request.POST.get('content')
        Item.objects.create(title=title, content=content)
        return redirect('item_list')
    return render(request, 'crud_app/item_form.html')

def item_delete(request, pk):
    item = get_object_or_404(Item, pk=pk)
    item.delete()
    return redirect('item_list')